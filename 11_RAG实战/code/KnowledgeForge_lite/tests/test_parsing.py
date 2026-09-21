"""11.2 解析与清洗：编码、表格、脚本、占位符。跑起来不调模型、不联网。"""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from forge_lite.data.ingest import (  # noqa: E402
    _read_docx,
    _read_html,
    _read_source,
    clean_text,
    decode_bytes,
    find_placeholders,
)


class DecodeTests(unittest.TestCase):
    """中文语料里 GBK 系文件至今常见，认不出编码等于整篇进不了库。"""

    def test_utf8_passes_through(self):
        self.assertEqual(decode_bytes("住宿 500 元".encode("utf-8")), "住宿 500 元")

    def test_gbk_falls_back_instead_of_crashing(self):
        """老系统导出的 txt 多半是 GBK，只试 UTF-8 会直接抛 UnicodeDecodeError。"""
        self.assertEqual(decode_bytes("住宿标准 500 元".encode("gb18030")), "住宿标准 500 元")

    def test_unknown_encoding_raises_with_a_readable_message(self):
        with self.assertRaises(ValueError) as ctx:
            decode_bytes(b"\xff\xfe\xff\xfe\xff\xfe")
        self.assertIn("编码", str(ctx.exception))


class CleanTextTests(unittest.TestCase):
    def test_control_characters_are_stripped_before_whitespace_collapse(self):
        """控制字符夹在词中间会让「同一句话」的哈希对不上：重传一次就变成"内容变了"。"""
        self.assertEqual(clean_text("住宿\x0b标准\x07为 500 元"), "住宿标准为 500 元")

    def test_full_width_and_nbspace_become_plain_spaces(self):
        """全角空格（U+3000）和不换行空格（U+00A0）在中文文档里遍地都是。"""
        self.assertEqual(clean_text("住宿\u3000标准\u00a0为 500 元"), "住宿 标准 为 500 元")

    def test_page_number_line_is_removed_but_body_numbers_survive(self):
        text = clean_text("第一页内容\n第 2 页\n金额为 500 元\nPage 3\n")
        self.assertNotIn("第 2 页", text)
        self.assertNotIn("Page 3", text)
        self.assertIn("500 元", text)

    def test_three_blank_lines_collapse_to_one_paragraph_break(self):
        self.assertEqual(clean_text("甲\n\n\n\n乙"), "甲\n\n乙")


class PlaceholderTests(unittest.TestCase):
    """未填模板是企业库里真实存在的坑：谁问都召回它，而它什么信息都没有。"""

    def test_unfilled_template_is_flagged(self):
        found = find_placeholders(clean_text("本制度适用于×××公司全体员工，具体标准待定。"))
        self.assertIn("叉号占位", found)
        self.assertIn("待填写", found)

    def test_filled_document_is_clean(self):
        self.assertEqual(find_placeholders(clean_text("本制度适用于全体员工，标准见附表。")), [])


class DocxTests(unittest.TestCase):
    """表格是制度文件里最关键的信息，丢了正文还通顺，只是答案永远查不到。"""

    def _make(self, tmp: Path) -> Path:
        from docx import Document as DocxDocument

        doc = DocxDocument()
        doc.add_paragraph("第一条 住宿标准")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "城市档位"
        table.cell(0, 1).text = "每人每天上限"
        table.cell(1, 0).text = "一线城市"
        table.cell(1, 1).text = "500 元"
        doc.add_paragraph("第二条 报销流程")
        path = tmp / "制度.docx"
        doc.save(str(path))
        return path

    def test_table_content_is_extracted(self):
        with tempfile.TemporaryDirectory() as tmp:
            text = _read_docx(self._make(Path(tmp)))
        self.assertIn("500 元", text, "表格里的金额必须读得出来")
        self.assertIn("一线城市", text, "金额要跟着它的档位一起出现")

    def test_paragraphs_and_tables_keep_document_order(self):
        """「下表说明标准」和那张表是连着的，拆开顺序就断了。"""
        with tempfile.TemporaryDirectory() as tmp:
            text = _read_docx(self._make(Path(tmp)))
        self.assertLess(text.index("第一条 住宿标准"), text.index("表格行:"))
        self.assertLess(text.index("表格行:"), text.index("第二条 报销流程"))

    def test_table_row_keeps_columns_on_one_line(self):
        """一行表格是一条完整记录，打散成单元格会让数字看不出属于谁。"""
        with tempfile.TemporaryDirectory() as tmp:
            text = _read_docx(self._make(Path(tmp)))
        self.assertIn("表格行: 一线城市 | 500 元", text)


class HtmlTests(unittest.TestCase):
    def _write(self, tmp: Path, body: str) -> Path:
        path = tmp / "页.html"
        path.write_text(f"<html><head><title>标题</title></head><body>{body}</body></html>",
                        encoding="utf-8")
        return path

    def test_script_and_style_are_dropped_whole(self):
        """脚本内容进了索引不只是噪声——里面的变量名还会被 BM25 当关键词召回。"""
        with tempfile.TemporaryDirectory() as tmp:
            text = _read_html(self._write(
                Path(tmp),
                "<style>.a{color:red}</style><script>var 内部代号 = '不该被检索到';</script>"
                "<p>住宿标准 500 元</p>",
            ))
        self.assertIn("住宿标准 500 元", text)
        self.assertNotIn("内部代号", text)
        self.assertNotIn("color:red", text)

    def test_table_cells_join_with_pipes(self):
        """和 docx 的表格行保持同一个形状，下游"表格不切碎"的规则才能一视同仁。"""
        with tempfile.TemporaryDirectory() as tmp:
            text = clean_text(_read_html(self._write(
                Path(tmp),
                "<table><tr><th>档位</th><th>上限</th></tr>"
                "<tr><td>一线</td><td>500 元</td></tr></table>",
            )))
        self.assertIn("档位 | 上限", text)
        self.assertIn("一线 | 500 元", text)

    def test_entities_are_unescaped_after_tag_stripping(self):
        """顺序不能反：先 unescape 的话，正文里写着的 &lt;p&gt; 会被当成真标签吃掉。"""
        with tempfile.TemporaryDirectory() as tmp:
            text = _read_html(self._write(Path(tmp), "<p>写法是 &lt;p&gt; 标签</p>"))
        self.assertIn("写法是 <p> 标签", text)


class SuffixRoutingTests(unittest.TestCase):
    def test_gbk_text_file_is_read_end_to_end(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "老制度.txt"
            path.write_bytes("住宿标准为每人每天不超过 500 元。".encode("gb18030"))
            self.assertIn("500 元", _read_source(path))

    def test_unknown_suffix_falls_back_to_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "随手记.log"
            path.write_text("住宿 500 元", encoding="utf-8")
            self.assertIn("500 元", _read_source(path))


if __name__ == "__main__":
    unittest.main()
